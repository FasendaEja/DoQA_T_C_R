"""Парсер DOCX файлов."""

from __future__ import annotations

from pathlib import Path

from .base import BaseParser, ParsedDocument


class DocxParser(BaseParser):
    """Парсер для .docx файлов."""

    @property
    def supported_extensions(self) -> list[str]:
        return [".docx"]

    def can_parse(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in self.supported_extensions

    def parse(self, file_path: Path) -> ParsedDocument:
        """Парсинг DOCX файла."""
        try:
            from docx import Document

            doc = Document(file_path)
            paragraphs = []
            fragments = []
            fragment_id = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

                    # Определяем тип параграфа по стилю
                    style_name = para.style.name if para.style else ""
                    if "Heading" in style_name:
                        fragment_id += 1
                        fragments.append({
                            "fragment_id": f"frag_{fragment_id}",
                            "title": text,
                            "text": text,
                            "type": "heading",
                            "level": self._get_heading_level(style_name),
                        })
                    else:
                        fragment_id += 1
                        fragments.append({
                            "fragment_id": f"frag_{fragment_id}",
                            "title": "",
                            "text": text,
                            "type": "paragraph",
                        })

            content = "\n\n".join(paragraphs)

            return ParsedDocument(
                source=str(file_path),
                filename=file_path.name,
                content=content,
                fragments=fragments,
                metadata={
                    "encoding": "utf-8",
                    "size": len(content),
                    "paragraphs": len(paragraphs),
                },
            )
        except ImportError:
            return ParsedDocument(
                source=str(file_path),
                filename=file_path.name,
                content="",
                errors=["python-docx не установлен. Установите: pip install python-docx"],
            )
        except Exception as e:
            return ParsedDocument(
                source=str(file_path),
                filename=file_path.name,
                content="",
                errors=[f"Ошибка парсинга DOCX: {e}"],
            )

    def parse_text(self, text: str, source: str = "text") -> ParsedDocument:
        """DOCX нельзя парсить из текста напрямую."""
        return ParsedDocument(
            source=source,
            filename="docx",
            content="",
            errors=["DOCX формат не поддерживает парсинг из текста"],
        )

    def _get_heading_level(self, style_name: str) -> int:
        """Извлечение уровня заголовка из имени стиля."""
        try:
            return int(style_name.replace("Heading", "").strip())
        except ValueError:
            return 1
